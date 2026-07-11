"""Simple Streamlit dashboard for the ISMS ISO Toolkit (hobby project).

Run with:
    streamlit run dashboard/app.py

Loads the committed example artifacts from examples/.
"""

import streamlit as st
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

st.set_page_config(page_title="ISMS Toolkit — Aether Labs (Hobby)", layout="wide")

ROOT = Path(__file__).resolve().parents[1]
RISK_XLSX = ROOT / "examples" / "risk-register-aether-labs.xlsx"
GAP_XLSX = ROOT / "examples" / "gap-analysis-aether-labs.xlsx"
POLICIES_DIR = ROOT / "policies"

st.title("ISMS ISO Toolkit — Hobby Explorer")
st.caption("Personal project for ISO 27001:2022 risk & compliance automation (fictional Aether Labs)")

with st.sidebar:
    st.header("About this hobby project")
    st.markdown("""
    Demonstrates generated artifacts for a fictional company.

    **Company:** Aether Labs (remote SaaS, ~22 people)  
    **Data source:** Committed examples in `examples/`

    All content is educational.
    """)
    st.info("To explore live data, run the generation scripts and reload.")

# Load functions
@st.cache_data
def load_risk_data():
    wb = load_workbook(RISK_XLSX, data_only=True)
    ws = wb["Risk Register"]
    data = []
    headers = [cell.value for cell in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            data.append(dict(zip(headers, row)))
    return pd.DataFrame(data), wb

@st.cache_data
def load_gap_data():
    wb = load_workbook(GAP_XLSX, data_only=True)
    ws = wb["Annex A Controls"]
    data = []
    headers = [cell.value for cell in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            data.append(dict(zip(headers, row)))
    return pd.DataFrame(data), wb

risk_df, risk_wb = load_risk_data()
gap_df, gap_wb = load_gap_data()

tab_risk, tab_heat, tab_gap, tab_pol = st.tabs(
    ["Risk Overview", "Heatmap & Top Risks", "Gap Analysis", "Policies"]
)

with tab_risk:
    st.subheader("Risk Register")
    st.dataframe(risk_df, use_container_width=True, height=420)
    col1, col2, col3 = st.columns(3)
    col1.metric("Total Risks", len(risk_df))
    if "Rating" in risk_df.columns:
        col2.metric("Critical Risks", len(risk_df[risk_df["Rating"] == "Critical"]))
    col3.metric("Data source", "examples/risk-register-aether-labs.xlsx")

with tab_heat:
    st.subheader("5×5 Heatmap + Top 5 Risks")
    st.markdown("The full colored heatmap is available in the generated Excel file.")
    st.write("**Top risks (by inherent score):**")
    if "Inherent Score" in risk_df.columns:
        top5 = risk_df.nlargest(5, "Inherent Score")[["ID", "Asset", "Threat", "Inherent Score", "Rating"]]
        st.dataframe(top5, use_container_width=True)
    else:
        st.dataframe(risk_df.head(5))

with tab_gap:
    st.subheader("Annex A (2022) Gap Analysis")
    st.dataframe(gap_df, use_container_width=True, height=420)
    st.caption("93 controls • Statuses assigned for fictional Aether Labs profile")

with tab_pol:
    st.subheader("Policy Templates")
    st.markdown("All policies are Markdown templates in the `policies/` folder with Annex A references.")
    policy_files = sorted(POLICIES_DIR.glob("*.md"))
    for p in policy_files:
        if p.name != "README.md":
            with st.expander(p.stem.replace("-", " ").title()):
                st.markdown(p.read_text()[:1500] + "\n\n_(truncated — see full file)_")

st.markdown("---")
st.caption("Hobby project — everything is fictional and for demonstration purposes only. Not for production use.")
