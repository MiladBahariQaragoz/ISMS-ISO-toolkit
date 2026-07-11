"""Structure and smoke tests for the gap analysis Excel generator.

Per plan P3-4.
"""

import json
from pathlib import Path

from openpyxl import load_workbook

from isms_toolkit.excel.gap_workbook import create_gap_workbook, save_workbook

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE_XLSX = ROOT / "examples" / "gap-analysis-aether-labs.xlsx"
CONTROLS_FILE = ROOT / "data" / "annex_a_controls.json"


def test_example_workbook_exists_and_has_expected_sheets():
    assert EXAMPLE_XLSX.exists(), "Run scripts/generate_gap_analysis.py first"
    wb = load_workbook(EXAMPLE_XLSX)
    sheets = wb.sheetnames
    assert "Document Control" in sheets
    assert "Annex A Controls" in sheets
    assert "Summary" in sheets
    assert "SoA Excerpt" in sheets


def test_annex_a_controls_has_93_rows_and_formatting():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Annex A Controls"]

    # Header + 93 controls
    assert ws.max_row == 94, f"Expected 94 rows, got {ws.max_row}"

    # Auto filter and freeze
    assert ws.auto_filter.ref is not None
    assert ws.freeze_panes is not None

    # Header styling
    first_cell = ws.cell(row=1, column=1)
    assert first_cell.font.bold

    # Status column has colors (check a few)
    for row in range(2, min(10, ws.max_row + 1)):
        status_cell = ws.cell(row=row, column=4)
        if status_cell.fill.fgColor.rgb and status_cell.fill.fgColor.rgb not in ("00000000", None):
            break
    else:
        assert False, "Expected at least one status cell with color fill"


def test_theme_distribution_in_controls():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Annex A Controls"]

    themes = {}
    for row in range(2, ws.max_row + 1):
        theme = ws.cell(row=row, column=2).value
        if theme:
            themes[theme] = themes.get(theme, 0) + 1

    assert themes.get("Organizational") == 37
    assert themes.get("People") == 8
    assert themes.get("Physical") == 14
    assert themes.get("Technological") == 34


def test_summary_sheet_math():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Summary"]

    # Total
    assert ws["B3"].value == 93

    # Check some counts exist (not exact since logic, but structure)
    found_org = False
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        if row[0] and "Organizational" in str(row[0]):
            found_org = True
            break
    assert found_org


def test_soa_excerpt_has_entries():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["SoA Excerpt"]

    # Title
    assert "Statement of Applicability" in str(ws["A1"].value)

    # Has header + at least some data rows
    assert ws.max_row >= 10

    # First data row has ID like A.5.x
    first_id = ws.cell(row=6, column=1).value
    assert first_id and first_id.startswith("A.")


def test_generator_produces_valid_workbook(tmp_path):
    """Regenerate and validate structure."""
    controls = json.load(CONTROLS_FILE.open())
    assert len(controls) == 93

    wb = create_gap_workbook(controls, company="Aether Labs")

    out = tmp_path / "test-gap.xlsx"
    save_workbook(wb, out)

    wb2 = load_workbook(out)
    assert "Annex A Controls" in wb2.sheetnames
    ws = wb2["Annex A Controls"]
    assert ws.max_row == 94
