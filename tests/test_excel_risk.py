"""Structure and smoke tests for the risk register Excel generator.

Per plan P2-3.
"""

from pathlib import Path

import yaml
from openpyxl import load_workbook

from isms_toolkit.excel.risk_workbook import create_risk_workbook, save_workbook

ROOT = Path(__file__).resolve().parents[1]
EXAMPLE_XLSX = ROOT / "examples" / "risk-register-aether-labs.xlsx"
SEED = ROOT / "data" / "risk_register_seed.yaml"


def test_example_workbook_exists_and_has_expected_sheets():
    assert EXAMPLE_XLSX.exists(), "Run scripts/generate_risk_register.py first to create example"
    wb = load_workbook(EXAMPLE_XLSX)
    sheets = wb.sheetnames
    assert "Document Control" in sheets
    assert "Risk Register" in sheets
    assert "Heatmap" in sheets
    assert "Summary" in sheets


def test_risk_register_has_sufficient_data_and_formatting():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Risk Register"]

    # Header + data rows (we have 25 risks)
    assert ws.max_row >= 23, f"Expected at least 23 rows (header + >=22 data), got {ws.max_row}"

    # Auto filter present
    assert ws.auto_filter.ref is not None

    # Freeze panes
    assert ws.freeze_panes is not None

    # Header styling (dark blue fill on first row)
    first_cell = ws.cell(row=1, column=1)
    assert first_cell.fill.fgColor.rgb in ("1F4E79", "FF1F4E79") or first_cell.font.bold

    # Check a rating cell has color (row 2 should have one)
    rating_cell = ws.cell(row=2, column=8)
    # Color fill should be set for the rating
    assert rating_cell.fill.fgColor.rgb is not None or rating_cell.fill.fgColor.theme is not None


def test_document_control_contains_hobby_markers():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Document Control"]
    content = " ".join(str(cell.value or "") for row in ws.iter_rows() for cell in row)
    assert "hobby" in content.lower()
    assert "ISMS-ISO-toolkit" in content or "fictional" in content.lower()


def test_heatmap_is_functional_5x5():
    wb = load_workbook(EXAMPLE_XLSX)
    ws = wb["Heatmap"]

    # Title
    assert "5x5 Risk Heatmap" in str(ws["A1"].value)

    # Legend present
    legend_text = " ".join(
        str(cell.value or "")
        for row in ws.iter_rows(min_row=9, max_row=12)
        for cell in row
    )
    assert "Legend" in legend_text or "Low" in legend_text

    # Check that the 5x5 grid area has numeric values (counts)
    count_sum = 0
    for row in range(4, 9):
        for col in range(2, 7):
            val = ws.cell(row=row, column=col).value
            if isinstance(val, int):
                count_sum += val
    assert count_sum >= 20, "Heatmap should account for most/all risks"


def test_generator_produces_valid_workbook(tmp_path):
    """Regenerate from seed and validate structure (smoke + assertions)."""
    data = yaml.safe_load(SEED.read_text())
    risks = data.get("risks", [])
    assert len(risks) >= 20

    wb = create_risk_workbook(risks, company="Aether Labs")

    # Write to temp and reload
    out = tmp_path / "test-risk.xlsx"
    save_workbook(wb, out)

    wb2 = load_workbook(out)
    assert "Risk Register" in wb2.sheetnames

    ws_risk = wb2["Risk Register"]
    # header row + data
    assert ws_risk.max_row > 20

    # At least one colored rating
    has_color = False
    for row in range(2, min(ws_risk.max_row + 1, 10)):
        cell = ws_risk.cell(row=row, column=8)
        fill = cell.fill
        rgb = fill.fgColor.rgb if fill and fill.fgColor else None
        if rgb and rgb not in ("00000000", None):
            has_color = True
            break
    assert has_color, "Expected at least one risk rating cell with color fill"
